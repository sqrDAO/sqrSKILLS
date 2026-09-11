"""Offline regression tests for the six skills reviewed on 2026-09-11.

Network calls and credentials are replaced with fixtures. These tests cover
script contracts and data fidelity; prompt routing remains ungated.
"""

import contextlib
import importlib.util
import io
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]


def load(name, path):
    spec = importlib.util.spec_from_file_location(name, ROOT / path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


SUMMARY = load("remaining_summary", "telegram-group-summary/scripts/fetch_messages.py")
SEND = load("remaining_send", "telegram-send/scripts/send.py")
GROUPS = load("remaining_groups", "telegram-send/scripts/list_groups.py")
CHATS = load("remaining_chats", "list-telegram-chats/scripts/list_chats.py")
PLACES = load("remaining_places", "nearby-places-search/scripts/search_places.py")
LUMA = load("remaining_luma", "luma-calendar/scripts/luma_calendar.py")
WORKBOOK = load("remaining_workbook", "business-model-to-market/scripts/build_gtm_workbook.py")


class TelegramSummaryTests(unittest.TestCase):
    def test_caption_and_channel_post_are_supported_without_mutating_polling(self):
        caption = {"message_id": 1, "chat": {"id": -7}, "caption": "photo", "date": 1700000000}
        self.assertEqual("photo", SUMMARY._extract_from_value(caption, -7, None)[0]["text"])
        response = io.BytesIO(json.dumps({"ok": True, "result": [
            {"channel_post": {"message_id": 2, "chat": {"id": -7}, "text": "post", "date": 1700000001}}
        ]}).encode())
        with patch.dict(os.environ, {"TELEGRAM_BOT_TOKEN": "fixture"}), \
                patch("urllib.request.urlopen", return_value=response) as request:
            rows = SUMMARY._messages_from_bot_api(-7, 10, None)
        self.assertEqual(["post"], [row["text"] for row in rows])
        self.assertNotIn("allowed_updates", request.call_args.args[0].full_url)

    def test_jsonl_state_and_limit_are_handled(self):
        msg = {"message_id": 1, "chat": {"id": -7}, "text": "hello", "date": 1700000000}
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, {"OPENCLAW_STATE_DIR": tmp}):
            Path(tmp, "history.jsonl").write_text(json.dumps(msg) + "\n", encoding="utf-8")
            self.assertEqual(1, len(SUMMARY._messages_from_openclaw_state(-7, 10, None)))
            self.assertEqual([], SUMMARY._messages_from_openclaw_state(-7, 0, None))
        with self.assertRaises(ValueError):
            SUMMARY.fetch_messages(-7, 0, None)

    def test_cutoff_is_utc_even_when_host_timezone_differs(self):
        now = SUMMARY.datetime.datetime(2026, 9, 11, 12, tzinfo=SUMMARY.datetime.timezone.utc)
        real_datetime = SUMMARY.datetime.datetime
        class FrozenDateTime(real_datetime):
            @classmethod
            def now(cls, tz=None):
                return now if tz else now.replace(tzinfo=None)
        with patch.object(SUMMARY.datetime, "datetime", FrozenDateTime):
            self.assertEqual(now.timestamp() - 86400, SUMMARY.cutoff_timestamp(24))


class TelegramSendAndListTests(unittest.TestCase):
    def test_code_spans_are_literal(self):
        self.assertEqual("<code>user_name_here</code>", SEND.markdown_to_html("`user_name_here`"))
        self.assertIn("a_b_c = 1", SEND.markdown_to_html("```python\na_b_c = 1\n```"))

    def test_group_and_chat_order_is_total(self):
        with patch.object(GROUPS, "_groups_from_nanobot_sessions", return_value=[{"chat_id": -2}, {"chat_id": -1}]), \
                patch.object(GROUPS, "_groups_from_openclaw_sessions", return_value=[]):
            self.assertEqual([-2, -1], [x["chat_id"] for x in GROUPS.list_groups(False)])
        with tempfile.TemporaryDirectory() as tmp, patch.dict(os.environ, {"NANOBOT_SESSIONS_DIR": tmp}):
            Path(tmp, "telegram_-2.jsonl").write_text(json.dumps({"_type": "metadata", "updated_at": "2026-09-11"}) + "\n", encoding="utf-8")
            Path(tmp, "telegram_-1.jsonl").write_text(json.dumps({"_type": "metadata", "updated_at": "2026-09-10"}) + "\n", encoding="utf-8")
            result = CHATS.list_chats(False)
        self.assertEqual([-2, -1], [x["chat_id"] for x in result["groups"]])
        self.assertEqual("2026-09-11", result["groups"][0]["last_updated"])


class NearbyPlacesTests(unittest.TestCase):
    def test_memory_coordinates_are_scoped_to_the_matching_section(self):
        with tempfile.TemporaryDirectory() as tmp, patch.object(PLACES, "WORKSPACE_DIR", tmp):
            Path(tmp, "MEMORY.md").write_text("# Hanoi\nCoordinates: 21.0285, 105.8542\n\n# London\nCoordinates: 51.5074, -0.1278\n", encoding="utf-8")
            self.assertEqual((51.5074, -0.1278), PLACES.lookup_coords_in_memory("London"))

    def test_free_text_query_and_api_error_are_truthful(self):
        self.assertIsNone(PLACES.get_place_type("vegan restaurant"))
        with patch.dict(PLACES.__dict__, {"GOOGLE_PLACES_API_KEY": "fixture"}), \
                patch.object(PLACES, "geocode_location", return_value=(0, 0)), \
                patch.object(PLACES, "_places_request", side_effect=OSError("quota")):
            result = PLACES.search_nearby_places("cafe", "fixture")
        self.assertFalse(result["successful"])
        self.assertIn("quota", result["error"])
        self.assertFalse(PLACES.search_nearby_places("cafe", "fixture", 0)["successful"])


class LumaTests(unittest.TestCase):
    def test_current_routes_and_payload_names_are_used(self):
        calls = []
        with patch.object(LUMA, "api_get", side_effect=lambda path, params=None: calls.append((path, params)) or {}):
            LUMA.cmd_get_self(None)
            LUMA.cmd_get_calendar(None)
            LUMA.cmd_get_event(type("Args", (), {"api_id": "evt-1"})())
            LUMA.cmd_get_guests(type("Args", (), {"event_api_id": "evt-1", "pagination_cursor": None})())
        self.assertEqual("/v1/users/get-self", calls[0][0])
        self.assertEqual("/v1/calendars/get", calls[1][0])
        self.assertEqual(("/v1/events/get", {"event_id": "evt-1"}), calls[2])
        self.assertEqual(("/v1/events/guests/list", {"event_id": "evt-1"}), calls[3])

    def test_transport_errors_are_json(self):
        with patch.dict(os.environ, {"LUMA_API_KEY": "fixture"}), \
                patch("urllib.request.urlopen", side_effect=OSError("offline")), \
                contextlib.redirect_stdout(io.StringIO()) as output:
            with self.assertRaises(SystemExit) as raised:
                LUMA.api_get("/v1/users/get-self")
        self.assertEqual(1, raised.exception.code)
        self.assertEqual("offline", json.loads(output.getvalue())["error"]["error"])


@unittest.skipUnless(WORKBOOK.HAVE_OPENPYXL, "openpyxl is an optional workbook dependency")
class WorkbookTests(unittest.TestCase):
    def test_false_answer_is_not_tbd(self):
        from openpyxl import Workbook
        wb = Workbook()
        WORKBOOK.pair(wb.active, 1, "Decision maker?", False)
        self.assertFalse(wb.active["B1"].value)

    def test_missing_scores_are_blank_and_unranked(self):
        from openpyxl import Workbook
        wb = Workbook()
        wb.remove(wb.active)
        WORKBOOK.tab_decision_matrix(wb, {"decision_matrix": {"factors": [{"name": "Pain", "weight": 100}], "ideas": [{"name": "Idea"}]}})
        ws = wb["Decision Matrix"]
        self.assertIsNone(ws["B6"].value)
        self.assertIn('IF(COUNT', ws["D6"].value)
        self.assertIn('IF(D6=""', ws["E6"].value)


if __name__ == "__main__":
    unittest.main()
