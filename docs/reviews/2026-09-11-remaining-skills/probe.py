#!/usr/bin/env python3
"""Offline review evidence, not a prompt gate or a regression-test suite.

Run from any directory. Optional --dependency-dir supplies an isolated openpyxl
installation. Every network call is mocked; all runtime state is synthetic.
"""
import argparse
import contextlib
import datetime
import importlib.util
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import time
from unittest.mock import patch
import urllib.error

sys.dont_write_bytecode = True
ROOT = Path(__file__).resolve().parents[3]


def load(name, relative):
    spec = importlib.util.spec_from_file_location(name, ROOT / relative)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--dependency-dir', type=Path)
    args = parser.parse_args()
    if args.dependency_dir:
        sys.path.insert(0, str(args.dependency_dir))
    evidence = {}
    with tempfile.TemporaryDirectory() as tmp, contextlib.ExitStack() as stack:
        root = Path(tmp)
        stack.enter_context(patch.dict(os.environ, {
            'OPENCLAW_STATE_DIR': tmp, 'OPENCLAW_WORKSPACE_DIR': tmp,
            'NANOBOT_SESSIONS_DIR': str(root / 'nanobot'),
            'NANOBOT_CONFIG': str(root / 'missing-config.json'),
            'TELEGRAM_BOT_TOKEN': 'fixture-only', 'LUMA_API_KEY': 'fixture-only',
            'GOOGLE_PLACES_API_KEY': 'fixture-only',
        }, clear=True))
        stack.enter_context(patch('urllib.request.urlopen', side_effect=AssertionError('Network forbidden')))
        stack.enter_context(contextlib.redirect_stderr(io.StringIO()))
        places = load('review_places', 'nearby-places-search/scripts/search_places.py')
        summary = load('review_summary', 'telegram-group-summary/scripts/fetch_messages.py')
        send = load('review_send', 'telegram-send/scripts/send.py')
        chats = load('review_chats', 'list-telegram-chats/scripts/list_chats.py')
        groups = load('review_groups', 'telegram-send/scripts/list_groups.py')
        luma = load('review_luma', 'luma-calendar/scripts/luma_calendar.py')
        workbook = load('review_workbook', 'business-model-to-market/scripts/build_gtm_workbook.py')

        (root / 'MEMORY.md').write_text(
            '# Hanoi\nCoordinates: 21.0285, 105.8542\n\n'
            '# London\nCoordinates: 51.5074, -0.1278\n', encoding='utf-8')
        evidence['places_wrong_city'] = {
            'query': 'London', 'expected': [51.5074, -0.1278],
            'actual': places.lookup_coords_in_memory('London')}
        evidence['places_query_types'] = {
            q: places.get_place_type(q) for q in ['bookstore', 'barber', 'vegan restaurant', 'coworking space']}
        with patch.object(places, 'geocode_location', return_value=(0, 0)), \
                patch.object(places, '_places_request', side_effect=urllib.error.URLError('fixture outage')):
            result = places.search_nearby_places('cafe', 'fixture')
            evidence['places_outage'] = {
                'successful': result['successful'], 'error': result['error']}
        with patch.object(places, 'geocode_location', return_value=(0, 0)), \
                patch.object(places, '_google_nearby_search', return_value=[]) as calls:
            result = places.search_nearby_places('cafe', 'fixture', 500)
            evidence['places_radius_honored'] = {
                'requested_meters': 500, 'attempted_meters': [calls.call_args.args[4]],
                'result_count': result['data']['total_count']}

        original_tz = os.environ.get('TZ')
        try:
            os.environ['TZ'] = 'Asia/Ho_Chi_Minh'
            time.tzset()
            now = datetime.datetime(2026, 9, 11, 12, tzinfo=datetime.timezone.utc)
            class FrozenDateTime(datetime.datetime):
                @classmethod
                def utcnow(cls):
                    return now.replace(tzinfo=None)
            with patch.object(summary.datetime, 'datetime', FrozenDateTime), \
                    patch.object(summary, 'fetch_messages', return_value={'messages': []}) as fetch, \
                    patch.object(sys, 'argv', ['fetch_messages.py', '-1001', '--since-hours', '24']), \
                    contextlib.redirect_stdout(io.StringIO()):
                summary.main()
            actual = fetch.call_args.args[2]
            expected = now.timestamp() - 86400
            evidence['summary_cutoff'] = {'timezone': os.environ['TZ'],
                                         'error_hours': (actual - expected) / 3600}
        finally:
            if original_tz is None:
                os.environ.pop('TZ', None)
            else:
                os.environ['TZ'] = original_tz
            time.tzset()

        message = {'message_id': 1, 'chat': {'id': -1001}, 'text': 'hello', 'date': 1700000000}
        (root / 'history.jsonl').write_text(json.dumps(message) + '\n', encoding='utf-8')
        evidence['summary_jsonl_messages'] = summary._messages_from_openclaw_state(-1001, 100, None)
        (root / 'history.json').write_text(json.dumps([message]), encoding='utf-8')
        evidence['summary_json_positive_control'] = len(summary._messages_from_openclaw_state(-1001, 100, None))
        caption = dict(message)
        caption['caption'] = caption.pop('text')
        evidence['summary_caption_messages'] = summary._extract_from_value(caption, -1001, None)
        evidence['summary_zero_limit_count'] = len(summary._messages_from_openclaw_state(-1001, 0, None))
        response = io.BytesIO(json.dumps({'ok': True, 'result': [{'channel_post': message}]}).encode())
        with patch('urllib.request.urlopen', return_value=response) as api:
            result = summary._messages_from_bot_api(-1001, 100, None)
            evidence['summary_polling'] = {
                'request_query': api.call_args.args[0].full_url.split('?', 1)[1],
                'channel_messages_returned': len(result)}

        evidence['send_markdown'] = {
            text: send.markdown_to_html(text)
            for text in ['`user_name_here`', '```python\na_b_c = 1\n```', '**bold**']}
        with patch.object(send, 'send_message', return_value={'message_id': 1}), \
                patch.object(sys, 'argv', ['send.py', '-1001', 'fixture']):
            out = io.StringIO()
            with contextlib.redirect_stdout(out):
                send.main()
            evidence['send_stdout'] = out.getvalue().strip()

        session_dir = root / 'nanobot'
        session_dir.mkdir()
        for chat_id in [-1001, -1002]:
            (session_dir / f'telegram_{chat_id}.jsonl').write_text(
                json.dumps({'_type': 'metadata', 'created_at': '2026-09-01T00:00:00',
                            'updated_at': '2026-09-11T12:00:00'}) + '\n', encoding='utf-8')
        entries = ['telegram_-1001.jsonl', 'telegram_-1002.jsonl']
        evidence['chat_order'] = []
        evidence['group_order'] = []
        for order in [entries, list(reversed(entries))]:
            with patch('os.listdir', return_value=order):
                evidence['chat_order'].append(chats.list_chats(False)['groups'])
                evidence['group_order'].append(groups.list_groups(False))

        with patch.object(luma, 'api_get', return_value={}) as api, \
                contextlib.redirect_stdout(io.StringIO()):
            luma.cmd_list_events(argparse.Namespace(after=None, before=None, pagination_cursor=None))
        evidence['luma_upcoming_request'] = {'path': api.call_args.args[0], 'params': api.call_args.args[1]}
        with patch('urllib.request.urlopen', side_effect=urllib.error.URLError('fixture outage')):
            try:
                with contextlib.redirect_stdout(io.StringIO()):
                    luma.api_get('/v1/users/get-self')
            except (Exception, SystemExit) as exc:
                evidence['luma_transport_error'] = type(exc).__name__

        if not workbook.HAVE_OPENPYXL:
            evidence['workbook'] = {'skipped': 'openpyxl unavailable; use --dependency-dir or install documented dependency'}
        else:
            from openpyxl import Workbook, load_workbook
            wb = Workbook()
            ws = wb.active
            workbook.pair(ws, 1, 'Decision maker?', False)
            evidence['workbook_false_value'] = ws['B1'].value
            wb = Workbook()
            wb.remove(wb.active)
            workbook.tab_decision_matrix(wb, {'decision_matrix': {
                'factors': [{'name': 'Pain', 'weight': 100}],
                'ideas': [{'name': 'Unscored'}]}})
            ws = wb['Decision Matrix']
            evidence['workbook_missing_score'] = {
                'raw_score': ws['B6'].value, 'weighted': ws['C6'].value, 'rank': ws['E6'].value}
            example = ROOT / 'business-model-to-market/assets/example-answers.json'
            evidence['workbook_builds'] = {}
            for flag, expected_count in [(None, 11), ('--canvas-only', 1),
                                         ('--partnership-only', 1), ('--ideation-only', 2)]:
                target = root / ((flag or 'all') + '.xlsx')
                argv = ['build_gtm_workbook.py', str(example), '-o', str(target)]
                if flag:
                    argv.append(flag)
                with patch.object(sys, 'argv', argv), contextlib.redirect_stdout(io.StringIO()) as out:
                    workbook.main()
                generated = load_workbook(target)
                evidence['workbook_builds'][flag or 'all'] = {
                    'sheets': len(generated.sheetnames), 'expected': expected_count,
                    'stdout_is_json': out.getvalue().lstrip().startswith('{')}
                generated.close()
    print(json.dumps(evidence, indent=2, ensure_ascii=False))


if __name__ == '__main__':
    main()
